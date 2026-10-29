"""Checks for the couple's pages: signing in with the authenticator code, the overview, the guest list,
one guest, Excel and WhatsApp."""
import secrets
import time
from io import BytesIO
from urllib.parse import unquote

from allauth.mfa.totp.internal.auth import TOTP, format_hotp_value, generate_totp_secret, hotp_value
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from django.utils.html import escape
from openpyxl import Workbook, load_workbook

from .dashboard import STAGES
from .models import Guest, ReplyTo
from .phones import normalise_phone

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def admin_with_two_step(name='couple'):
    user = User.objects.create_user(name, password=secrets.token_urlsafe(), is_staff=True, is_superuser=True)
    TOTP.activate(user, generate_totp_secret())
    return user


def one_at_every_stage():
    now = timezone.now()
    return {
        'unsent': Guest.objects.create(name='Nimal', phone='94770000001'),
        'sent': Guest.objects.create(name='Amali', phone='94770000002', sent_at=now),
        'opened': Guest.objects.create(name='Jayasinghe', seats=4, sent_at=now, first_opened_at=now, times_opened=2),
        'coming': Guest.objects.create(name='Sunil & Kamala', seats=2, sent_at=now, first_opened_at=now, reply='yes',
                                       people=2, replied_at=now, message='See you there!', side='bride'),
        'declined': Guest.objects.create(name='Bandara', sent_at=now, first_opened_at=now, reply='no', replied_at=now, side='groom'),
    }


class PhoneTests(TestCase):
    def test_every_way_of_writing_a_sri_lankan_number(self):
        for typed in ['077 128 1561', '0771281561', '771281561', '+94 77 128 1561', '0094771281561',
                      '+940771281561', '94771281561', '(077) 128-1561']:
            self.assertEqual(normalise_phone(typed), '94771281561', typed)

    def test_other_countries_keep_their_code(self):
        self.assertEqual(normalise_phone('+44 7911 123456'), '447911123456')

    def test_not_a_number(self):
        self.assertEqual(normalise_phone(''), '')
        for typed in ['call me', '12345', '077 128 15611111111', '0712']:
            with self.assertRaises(ValueError, msg=typed):
                normalise_phone(typed)


class StageTests(TestCase):
    def test_every_guest_is_at_exactly_one_stage(self):
        for stage, guest in one_at_every_stage().items():
            self.assertEqual(guest.stage, stage)
            self.assertEqual([key for key, _, where in STAGES if Guest.objects.filter(where, pk=guest.pk).exists()], [stage])


class SignInTests(TestCase):
    def setUp(self):
        self.password = secrets.token_urlsafe()
        self.user = User.objects.create_user('couple', password=self.password, is_staff=True, is_superuser=True)

    def test_the_pages_need_signing_in(self):
        for url in ['/admin/', '/admin/guests/', reverse('couple:export'), reverse('couple:import')]:
            response = self.client.get(url)
            self.assertEqual(response.status_code, 302, url)
            self.assertTrue(response['Location'].startswith(reverse('account_login')), url)
        # Django's own admin, and its password-only sign-in, are gone.
        self.assertEqual(self.client.get('/admin/login/').status_code, 404)

    def test_nobody_can_sign_up(self):
        response = self.client.post(reverse('account_signup'),
                                    {'username': 'someone', 'password1': 'x' * 20, 'password2': 'x' * 20})
        self.assertFalse(User.objects.filter(username='someone').exists())
        self.assertContains(response, 'closed')

    def test_an_account_without_an_authenticator_must_add_one_first(self):
        self.client.force_login(self.user)
        response = self.client.get('/admin/guests/')
        self.assertRedirects(response, reverse('mfa_activate_totp'), fetch_redirect_response=False)

    def test_password_and_code_together(self):
        secret = generate_totp_secret()
        TOTP.activate(self.user, secret)
        response = self.client.post(reverse('account_login'), {'login': 'couple', 'password': self.password})
        self.assertRedirects(response, reverse('mfa_authenticate'), fetch_redirect_response=False)
        # The password alone has not signed anyone in.
        self.assertEqual(self.client.get('/admin/').status_code, 302)

        now = int(time.time()) // 30
        right = format_hotp_value(hotp_value(secret, now))
        nearby = {format_hotp_value(hotp_value(secret, now + step)) for step in (-1, 0, 1)}
        wrong = next(code for code in (f'{n:06}' for n in range(1000000)) if code not in nearby)
        self.assertEqual(self.client.post(reverse('mfa_authenticate'), {'code': wrong}).status_code, 200)
        self.assertEqual(self.client.get('/admin/').status_code, 302)

        self.assertEqual(self.client.post(reverse('mfa_authenticate'), {'code': right}).status_code, 302)
        self.assertEqual(self.client.get('/admin/').status_code, 200)

    def test_a_signed_in_account_that_is_not_an_admin_gets_nothing(self):
        user = User.objects.create_user('friend', password=secrets.token_urlsafe())
        TOTP.activate(user, generate_totp_secret())
        self.client.force_login(user)
        self.assertEqual(self.client.get('/admin/guests/').status_code, 403)


class OverviewTests(TestCase):
    def setUp(self):
        self.client.force_login(admin_with_two_step())

    def test_the_numbers(self):
        one_at_every_stage()
        response = self.client.get('/admin/')
        counts = response.context['counts']
        self.assertEqual((counts['invited'], counts['seats'], counts['people']), (5, 9, 2))
        self.assertEqual([counts[key] for key, _, _ in STAGES], [1, 1, 1, 1, 1])
        self.assertEqual([(tile['key'], tile['value']) for tile in response.context['tiles']],
                         [('coming', 2), ('declined', 1), ('waiting', 2), ('unsent', 1)])
        sides = {side['label']: (side['invited'], side['people'], side['declined'], side['waiting']) for side in response.context['sides']}
        self.assertEqual(sides["Bride's side"], (1, 2, 0, 0))
        self.assertEqual(sides["Groom's side"], (1, 0, 1, 0))
        self.assertEqual(sides['No side set'], (3, 0, 0, 3))
        self.assertContains(response, 'See you there!')

    def test_the_garden_has_a_flower_for_every_invitation(self):
        one_at_every_stage()
        response = self.client.get('/admin/')
        self.assertEqual([flower['stage'] for flower in response.context['garden']], ['coming', 'declined', 'opened', 'sent', 'unsent'])
        self.assertContains(response, 'class="flower is-coming"')
        self.assertContains(response, 'coming, 2 of 2')

    def test_the_garden_live(self):
        guests = one_at_every_stage()
        flowers = self.client.get(reverse('couple:garden')).json()['flowers']
        self.assertEqual({flower['pk']: flower['stage'] for flower in flowers}[guests['opened'].pk], 'opened')
        Guest.objects.filter(pk=guests['opened'].pk).update(reply='yes', people=3)
        flowers = self.client.get(reverse('couple:garden')).json()['flowers']
        self.assertEqual({flower['pk']: flower['stage'] for flower in flowers}[guests['opened'].pk], 'coming')
        self.client.logout()
        self.assertEqual(self.client.get(reverse('couple:garden')).status_code, 302)

    def test_an_empty_list(self):
        response = self.client.get('/admin/')
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'No replies yet')


class GuestListTests(TestCase):
    def setUp(self):
        self.client.force_login(admin_with_two_step())
        self.guests = one_at_every_stage()
        self.guest = self.guests['unsent']

    def test_the_list_shows_each_guest_with_what_to_do(self):
        response = self.client.get('/admin/guests/')
        for guest in self.guests.values():
            self.assertContains(response, escape(guest.name))
            self.assertContains(response, f'data-link="/{guest.code}"')
        self.assertContains(response, reverse('couple:send', args=[self.guest.pk]))

    def test_the_steps_filter_the_list(self):
        response = self.client.get('/admin/guests/?show=coming')
        self.assertEqual([guest.name for guest in response.context['guests']], ['Sunil & Kamala'])
        response = self.client.get('/admin/guests/?q=jaya')
        self.assertEqual([guest.name for guest in response.context['guests']], ['Jayasinghe'])
        response = self.client.get('/admin/guests/?side=groom')
        self.assertEqual([guest.name for guest in response.context['guests']], ['Bandara'])
        response = self.client.get('/admin/guests/?show=waiting')
        self.assertEqual([guest.name for guest in response.context['guests']], ['Amali', 'Jayasinghe'])

    def test_finding_a_number_however_it_is_typed(self):
        for typed in ['077 000 0002', '0770000002', '+94 77 000 0002', '0002']:
            response = self.client.get('/admin/guests/', {'q': typed})
            self.assertEqual([guest.name for guest in response.context['guests']], ['Amali'], typed)

    def test_adding_a_guest_tidies_their_number(self):
        response = self.client.post(reverse('couple:guest_new'), {'name': 'Ruwan', 'name_si': '', 'phone': '077 765 4321',
                                                                  'seats': 2, 'side': 'bride', 'notes': ''})
        guest = Guest.objects.get(name='Ruwan')
        self.assertRedirects(response, reverse('couple:guest', args=[guest.pk]), fetch_redirect_response=False)
        self.assertEqual(guest.phone, '94777654321')

    def test_a_number_that_is_not_one_is_refused(self):
        response = self.client.post(reverse('couple:guest_new'), {'name': 'Ruwan', 'phone': 'ask Nimal', 'seats': 1})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'does not look like a phone number')
        self.assertFalse(Guest.objects.filter(name='Ruwan').exists())

    def test_changing_a_guest(self):
        url = reverse('couple:guest', args=[self.guest.pk])
        self.assertContains(self.client.get(url), f'/{self.guest.code}')
        self.client.post(url, {'name': 'Nimal Perera', 'name_si': 'නිමල්', 'phone': '0770000001', 'seats': 2, 'side': 'groom', 'notes': ''})
        self.guest.refresh_from_db()
        self.assertEqual((self.guest.name, self.guest.name_si, self.guest.seats), ('Nimal Perera', 'නිමල්', 2))

    def test_taking_a_guest_off_the_list(self):
        url = reverse('couple:guest_delete', args=[self.guest.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        self.client.post(url)
        self.assertFalse(Guest.objects.filter(pk=self.guest.pk).exists())

    def test_whatsapp_opens_with_their_link_and_marks_it_sent(self):
        url = reverse('couple:send', args=[self.guest.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        response = self.client.post(url)
        self.assertTrue(response['Location'].startswith('https://wa.me/94770000001?text='))
        text = unquote(response['Location'])
        self.assertIn('Dear Nimal,', text)
        self.assertIn(f'http://testserver/{self.guest.code}', text)
        self.guest.refresh_from_db()
        first = self.guest.sent_at
        self.assertIsNotNone(first)
        self.client.post(url)
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.sent_at, first)

    def test_a_reminder(self):
        guest = self.guests['sent']
        sent = guest.sent_at
        response = self.client.post(reverse('couple:send', args=[guest.pk]), {'kind': 'reminder'})
        self.assertIn('gentle reminder', unquote(response['Location']))
        guest.refresh_from_db()
        self.assertIsNotNone(guest.reminded_at)
        self.assertEqual(guest.sent_at, sent)

    @override_settings(SITE_URL='https://naveenanddulanjani.com')
    def test_links_use_the_real_address_once_it_is_set(self):
        response = self.client.post(reverse('couple:send', args=[self.guest.pk]))
        self.assertIn(f'https://naveenanddulanjani.com/{self.guest.code}', unquote(response['Location']))

    def test_sending_again_after_they_replied(self):
        guest = self.guests['coming']
        guest.phone = '94770000007'
        guest.save()
        self.assertContains(self.client.get(reverse('couple:guest', args=[guest.pk])), 'Send the invitation again')
        response = self.client.post(reverse('couple:send', args=[guest.pk]))
        self.assertTrue(response['Location'].startswith('https://wa.me/94770000007?text='))
        guest.refresh_from_db()
        self.assertEqual((guest.stage, guest.people), ('coming', 2))

    def test_no_whatsapp_without_a_number(self):
        guest = self.guests['opened']
        response = self.client.post(reverse('couple:send', args=[guest.pk]))
        self.assertRedirects(response, reverse('couple:guest', args=[guest.pk]), fetch_redirect_response=False)
        self.assertContains(self.client.get('/admin/guests/'), 'Add their number')

    def test_recording_a_reply_given_by_phone(self):
        guest = self.guests['opened']
        url = reverse('couple:guest_fix', args=[guest.pk])
        self.assertEqual(self.client.get(url).status_code, 405)
        self.client.post(url, {'do': 'coming', 'people': 5})  # more than the four seats kept
        guest.refresh_from_db()
        self.assertEqual(guest.reply, '')
        self.client.post(url, {'do': 'coming', 'people': 3})
        guest.refresh_from_db()
        self.assertEqual((guest.stage, guest.people), ('coming', 3))
        self.assertIsNotNone(guest.replied_at)
        self.client.post(url, {'do': 'not-coming'})
        guest.refresh_from_db()
        self.assertEqual((guest.stage, guest.people), ('declined', 0))

    def test_clearing_a_reply_lets_them_answer_again(self):
        guest = self.guests['coming']
        self.client.post(reverse('couple:guest_fix', args=[guest.pk]), {'do': 'clear'})
        guest.refresh_from_db()
        self.assertEqual((guest.stage, guest.people, guest.message, guest.replied_at), ('opened', 0, '', None))

    def test_marking_as_sent_and_not_sent(self):
        url = reverse('couple:guest_fix', args=[self.guest.pk])
        self.client.post(url, {'do': 'sent'})
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.stage, 'sent')
        self.client.post(url, {'do': 'unsent'})
        self.guest.refresh_from_db()
        self.assertEqual((self.guest.stage, self.guest.reminded_at), ('unsent', None))

    def test_the_help_page(self):
        self.assertContains(self.client.get(reverse('couple:help')), 'A guest phoned or told me in person')

    def test_the_couple_looking_at_a_guest_page_does_not_count_as_opening(self):
        self.client.post(f'/{self.guest.code}/opened')
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.times_opened, 0)


class ReplyNumberTests(TestCase):
    def test_it_needs_signing_in(self):
        self.assertEqual(self.client.get(reverse('couple:numbers')).status_code, 302)

    def test_numbers_are_tidied_and_saved(self):
        self.client.force_login(admin_with_two_step())
        self.assertContains(self.client.get(reverse('couple:numbers')), ReplyTo.MAIN)
        response = self.client.post(reverse('couple:numbers'), {'bride': '077 000 0101', 'groom': '+94 77 000 0202', 'everyone': '0771281561'})
        self.assertRedirects(response, reverse('couple:numbers'), fetch_redirect_response=False)
        numbers = ReplyTo.current()
        self.assertEqual((numbers.bride, numbers.groom, numbers.everyone), ('94770000101', '94770000202', '94771281561'))
        self.assertContains(self.client.get('/admin/'), 'Replies go to +94770000101')

    def test_home_asks_for_the_numbers_until_both_are_set(self):
        self.client.force_login(admin_with_two_step())
        self.assertContains(self.client.get('/admin/'), "Add the bride's and the groom's WhatsApp number")
        self.client.post(reverse('couple:numbers'), {'bride': '0770000101', 'groom': '', 'everyone': '0771281561'})
        self.assertContains(self.client.get('/admin/'), "Add the groom's WhatsApp number")
        self.client.post(reverse('couple:numbers'), {'bride': '0770000101', 'groom': '0770000202', 'everyone': '0771281561'})
        self.assertNotContains(self.client.get('/admin/'), 'Add the numbers')
        self.assertContains(self.client.get('/admin/'), 'Reply numbers')

    def test_a_wrong_number_or_no_main_number_is_refused(self):
        self.client.force_login(admin_with_two_step())
        response = self.client.post(reverse('couple:numbers'), {'bride': 'ask Malini', 'groom': '', 'everyone': ''})
        self.assertContains(response, 'does not look like a phone number')
        self.assertContains(response, 'Keep a main number')
        self.assertEqual(ReplyTo.current().bride, '')


class VercelAddressTests(TestCase):
    @override_settings(VERCEL_OWN_HOST='wedding-invite-3f8iv4ffm-mark-6689.vercel.app', VERCEL_MAIN_HOST='wedding-invite-mark-6689.vercel.app')
    def test_a_deployment_address_goes_to_the_main_one(self):
        response = self.client.get('/admin/guests/?show=coming', HTTP_HOST='wedding-invite-3f8iv4ffm-mark-6689.vercel.app')
        self.assertEqual(response['Location'], 'https://wedding-invite-mark-6689.vercel.app/admin/guests/?show=coming')
        response = self.client.get('/', HTTP_HOST='wedding-invite-mark-6689.vercel.app')
        self.assertEqual(response.status_code, 200)

    def test_off_vercel_nothing_moves(self):
        self.assertEqual(self.client.get('/', HTTP_HOST='localhost').status_code, 200)


class ExcelTests(TestCase):
    def setUp(self):
        self.client.force_login(admin_with_two_step())

    @staticmethod
    def sheet(*rows):
        book = Workbook()
        for row in rows:
            book.active.append(row)
        out = BytesIO()
        book.save(out)
        return SimpleUploadedFile('guests.xlsx', out.getvalue(), content_type=XLSX)

    def test_import_shows_what_will_happen_then_does_it(self):
        nimal = Guest.objects.create(name='Nimal Perera', phone='94771111111', notes='from the office')
        upload = self.sheet(
            ['Wedding guest list'],
            ['Guest Name', 'සිංහල නම', 'Mobile', 'Pax', 'Side', 'Remarks', 'Table'],
            ['Nimal Perera', 'නිමල්', '077 111 1111', 2, 'Groom', None, 5],
            ['Sunil & Kamala', None, '+94 77 222 2222', 2, "Bride's side", 'cousins', None],
            ['Ruwan', None, 'not a number', 1, None, None, None],
            ['Sunil & Kamala', None, '0772222222', 2, None, None, None],
            [None, None, None, None, None, None, None],
            ['Amara', None, None, 'two', None, None, None],
        )
        url = reverse('couple:import')
        response = self.client.post(url, {'sheet': upload})
        self.assertEqual((response.context['adding'], response.context['updating'], response.context['skipping']), (1, 1, 3))
        self.assertContains(response, 'Ruwan')  # a row left out still says whose it was
        self.assertEqual(Guest.objects.count(), 1)  # nothing changes before going ahead

        response = self.client.post(url, {'confirm': '1'})
        self.assertRedirects(response, reverse('couple:guests'), fetch_redirect_response=False)
        nimal.refresh_from_db()
        self.assertEqual((nimal.name_si, nimal.seats, nimal.side, nimal.notes), ('නිමල්', 2, 'groom', 'from the office'))
        added = Guest.objects.get(name='Sunil & Kamala')
        self.assertEqual((added.phone, added.seats, added.side, added.notes), ('94772222222', 2, 'bride', 'cousins'))
        self.assertEqual(Guest.objects.count(), 2)

    def test_a_sheet_without_a_name_column(self):
        response = self.client.post(reverse('couple:import'), {'sheet': self.sheet(['Phone'], ['0771111111'])})
        self.assertContains(response, 'No &quot;Name&quot; column')

    def test_a_file_that_is_not_excel(self):
        upload = SimpleUploadedFile('guests.xlsx', b'not a workbook', content_type=XLSX)
        self.assertContains(self.client.post(reverse('couple:import'), {'sheet': upload}), 'not an Excel')

    def test_download_has_every_guest_and_their_link(self):
        guest = Guest.objects.create(name='Nimal', phone='94771234567', reply='yes', people=1)
        response = self.client.get(reverse('couple:export'))
        self.assertEqual(response['Content-Type'], XLSX)
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[0][:9], ('Name', 'Sinhala name', 'Phone', 'Seats', 'Side', 'Notes', 'Link', 'Sent', 'Last reminded'))
        self.assertEqual(rows[1][0], 'Nimal')
        self.assertEqual(rows[1][6], f'http://testserver/{guest.code}')
        self.assertEqual(rows[1][11], 'Coming')

    def test_the_empty_list_to_fill_in(self):
        response = self.client.get(reverse('couple:blank'))
        first = next(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(first, ('Name', 'Sinhala name', 'Phone', 'Seats', 'Side', 'Notes'))
