"""The guest list in and out of Excel.

Reading is forgiving about the sheet: the first row with a name column is the heading row, columns
can be in any order and called most things a person would call them, and anything else is ignored.
"""
import re
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from .models import Guest
from .phones import normalise_phone

COLUMNS = {
    'name': ['name', 'guest', 'guest name', 'guests name', 'full name', 'invitation name', 'invite name'],
    'name_si': ['sinhala name', 'name in sinhala', 'name sinhala', 'sinhala', 'සිංහල නම'],
    'phone': ['phone', 'phone number', 'mobile', 'mobile number', 'whatsapp', 'whatsapp number', 'number', 'contact', 'tel'],
    'seats': ['seats', 'seat', 'pax', 'people', 'guests', 'count', 'number of guests', 'no of guests', 'invited'],
    'side': ['side', 'bride or groom', 'from'],
    'notes': ['notes', 'note', 'remarks', 'remark', 'comment', 'comments'],
}
HEADINGS = ['Name', 'Sinhala name', 'Phone', 'Seats', 'Side', 'Notes']
MAX_ROWS = 2000


def _heading(value):
    # Punctuation only: the Sinhala vowel signs are not "word" characters to a regular expression.
    return ' '.join(re.sub(r"[.,:;()'’\"/\\_*#?-]", ' ', str(value or '').lower()).split())


def _side(value):
    value = str(value or '').strip().lower()
    if not value:
        return ''
    if value.startswith('b') or 'bride' in value:
        return 'bride'
    if value.startswith('g') or 'groom' in value:
        return 'groom'
    raise ValueError


def _text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read(file):
    """What importing this sheet would do, row by row, without changing anything.

    Returns (rows, columns): each row is {'line', 'action', 'id', 'fields', 'reason'}, where action is
    'add', 'update' or 'skip'; columns are the headings that were used.
    """
    try:
        sheet = load_workbook(file, read_only=True, data_only=True).active
    except Exception as error:  # openpyxl raises many kinds for a file that is not a workbook
        raise ValueError('This is not an Excel (.xlsx) file.') from error

    lines = sheet.iter_rows(values_only=True)
    places, used = {}, []
    for number, cells in enumerate(lines, start=1):
        headings = [_heading(cell) for cell in cells]
        if any(heading in COLUMNS['name'] for heading in headings):
            for field, names in COLUMNS.items():
                for at, heading in enumerate(headings):
                    if heading in names and field not in places:
                        places[field] = at
                        used.append(str(cells[at]).strip())
            break
        if number >= 20:
            break
    if 'name' not in places:
        raise ValueError('No "Name" column was found in the first rows of the sheet.')

    by_phone = {guest.phone: guest for guest in Guest.objects.exclude(phone='')}
    by_name = {guest.name.casefold(): guest for guest in Guest.objects.all()}
    seen, rows = set(), []
    for number, cells in enumerate(lines, start=number + 1):
        cell = lambda field: _text(cells[places[field]]) if field in places and places[field] < len(cells) else ''
        name = cell('name')
        if not name:
            continue
        if len(rows) >= MAX_ROWS:
            raise ValueError(f'The sheet has more than {MAX_ROWS} guests.')
        # The name is kept even on a row that is left out, so the preview says whose row it was.
        row = {'line': number, 'action': 'skip', 'id': None, 'fields': {'name': name[:100]}, 'reason': ''}
        rows.append(row)
        try:
            phone = normalise_phone(cell('phone'))
        except ValueError:
            row['reason'] = f'"{cell("phone")}" is not a phone number'
            continue
        seats = cell('seats')
        if seats and not (seats.isdigit() and 1 <= int(seats) <= 50):
            row['reason'] = f'"{seats}" is not a number of seats'
            continue
        try:
            side = _side(cell('side'))
        except ValueError:
            row['reason'] = f'"{cell("side")}" is not bride or groom'
            continue
        key = phone or name.casefold()
        if key in seen:
            row['reason'] = 'appears twice in the sheet'
            continue
        seen.add(key)
        # Empty cells leave what is already there alone, so a Sinhala name typed in the admin pages
        # survives a later import of a sheet without one.
        fields = {'name': name[:100], 'name_si': cell('name_si')[:100], 'phone': phone,
                  'seats': int(seats) if seats else None, 'side': side, 'notes': cell('notes')}
        row['fields'] = {field: value for field, value in fields.items() if value not in ('', None)}
        existing = by_phone.get(phone) if phone else None
        if existing is None:
            same_name = by_name.get(name.casefold())
            # Someone of the same name is the same guest, unless both have numbers and they differ.
            if same_name and not (phone and same_name.phone):
                existing = same_name
        if existing:
            row.update(action='update', id=existing.pk)
        else:
            row['action'] = 'add'
    return rows, used


def apply(rows):
    """Carry out what read() found. Returns how many guests were added and updated."""
    added = updated = 0
    for row in rows:
        if row['action'] == 'add':
            Guest.objects.create(**row['fields'])
            added += 1
        elif row['action'] == 'update':
            guest = Guest.objects.filter(pk=row['id']).first()
            if guest:
                for field, value in row['fields'].items():
                    setattr(guest, field, value)
                guest.save()
                updated += 1
    return added, updated


def _book(headings):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Guests'
    sheet.append(headings)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = 'A2'
    return book, sheet


def _bytes(book):
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def blank():
    book, sheet = _book(HEADINGS)
    for column, width in zip('ABCDEF', [30, 30, 18, 8, 10, 40]):
        sheet.column_dimensions[column].width = width
    return _bytes(book)


def write(guests, link_of):
    """Every guest, with their link, whether it was sent and opened, and their reply."""
    book, sheet = _book(HEADINGS + ['Link', 'Sent', 'Last reminded', 'First opened', 'Times opened', 'Reply',
                                    'People coming', 'Their note', 'Replied'])
    local = lambda moment: timezone.localtime(moment).replace(tzinfo=None) if moment else None
    for guest in guests:
        sheet.append([guest.name, guest.name_si, guest.phone, guest.seats, guest.get_side_display() if guest.side else '',
                      guest.notes, link_of(guest), local(guest.sent_at), local(guest.reminded_at),
                      local(guest.first_opened_at), guest.times_opened, guest.get_reply_display() if guest.reply else '',
                      guest.people if guest.reply == 'yes' else None, guest.message, local(guest.replied_at)])
    for column, width in zip('ABCDEFGHIJKLMNO', [30, 30, 16, 7, 14, 30, 34, 17, 17, 17, 8, 12, 9, 40, 17]):
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in (row[7], row[8], row[9], row[14]):
            cell.number_format = 'd mmm yyyy h:mm'
    return _bytes(book)
